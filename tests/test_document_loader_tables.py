from pathlib import Path
import zipfile

from docx import Document

from app.document_loader import load_document, load_document_content


def test_hwpx_preserves_table_rows(tmp_path: Path):
    path = tmp_path / "sample.hwpx"
    xml = """<?xml version="1.0" encoding="UTF-8"?>
    <hs:sec xmlns:hs="urn:section" xmlns:hp="urn:para">
      <hp:p><hp:run><hp:t>사업 운영 계획</hp:t></hp:run></hp:p>
      <hp:tbl>
        <hp:tr>
          <hp:tc><hp:subList><hp:p><hp:run><hp:t>항목</hp:t></hp:run></hp:p></hp:subList></hp:tc>
          <hp:tc><hp:subList><hp:p><hp:run><hp:t>세부 내용</hp:t></hp:run></hp:p></hp:subList></hp:tc>
        </hp:tr>
        <hp:tr>
          <hp:tc><hp:subList><hp:p><hp:run><hp:t>모니터링</hp:t></hp:run></hp:p></hp:subList></hp:tc>
          <hp:tc>
            <hp:subList>
              <hp:p><hp:run><hp:t>정기 점검</hp:t></hp:run></hp:p>
              <hp:p><hp:run><hp:t>결과 피드백</hp:t></hp:run></hp:p>
            </hp:subList>
          </hp:tc>
        </hp:tr>
      </hp:tbl>
      <hp:p><hp:run><hp:t>이상 끝.</hp:t></hp:run></hp:p>
    </hs:sec>
    """

    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("Contents/section0.xml", xml)

    text = load_document(path)

    assert "사업 운영 계획" in text
    assert "항목 | 세부 내용" in text
    assert "모니터링 | 정기 점검 / 결과 피드백" in text
    assert text.count("모니터링") == 1
    assert "이상 끝." in text


def test_docx_preserves_body_table_order(tmp_path: Path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("표 앞 문단")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "내용"
    table.cell(1, 0).text = "피드백"
    table.cell(1, 1).text = "의견 반영"
    doc.add_paragraph("표 뒤 문단")
    doc.save(path)

    text = load_document(path)
    lines = text.splitlines()

    assert lines == [
        "표 앞 문단",
        "항목 | 내용",
        "피드백 | 의견 반영",
        "표 뒤 문단",
    ]


def test_hwpx_reconstructs_merged_cell_columns(tmp_path: Path):
    path = tmp_path / "merged.hwpx"
    xml = """<?xml version="1.0" encoding="UTF-8"?>
    <hs:sec xmlns:hs="urn:section" xmlns:hp="urn:para">
      <hp:p>
        <hp:run>
          <hp:tbl colCnt="3" rowCnt="2">
            <hp:tr>
              <hp:tc>
                <hp:subList><hp:p><hp:run><hp:t>구분</hp:t></hp:run></hp:p></hp:subList>
                <hp:cellAddr colAddr="0" rowAddr="0"/>
                <hp:cellSpan colSpan="1" rowSpan="1"/>
              </hp:tc>
              <hp:tc>
                <hp:subList><hp:p><hp:run><hp:t>운영 내용</hp:t></hp:run></hp:p></hp:subList>
                <hp:cellAddr colAddr="1" rowAddr="0"/>
                <hp:cellSpan colSpan="2" rowSpan="1"/>
              </hp:tc>
            </hp:tr>
            <hp:tr>
              <hp:tc>
                <hp:subList><hp:p><hp:run><hp:t>1</hp:t></hp:run></hp:p></hp:subList>
                <hp:cellAddr colAddr="0" rowAddr="1"/>
              </hp:tc>
              <hp:tc>
                <hp:subList><hp:p><hp:run><hp:t>피드백</hp:t></hp:run></hp:p></hp:subList>
                <hp:cellAddr colAddr="1" rowAddr="1"/>
              </hp:tc>
              <hp:tc>
                <hp:subList><hp:p><hp:run><hp:t>모니터링</hp:t></hp:run></hp:p></hp:subList>
                <hp:cellAddr colAddr="2" rowAddr="1"/>
              </hp:tc>
            </hp:tr>
          </hp:tbl>
        </hp:run>
      </hp:p>
    </hs:sec>
    """

    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("Contents/section0.xml", xml)

    text = load_document(path)

    assert "구분 | 운영 내용" in text
    assert "1 | 피드백 | 모니터링" in text
    assert text.count("1 | 피드백 | 모니터링") == 1


def test_hwp_uses_local_hancom_conversion(tmp_path: Path, monkeypatch):
    source = tmp_path / "sample.hwp"
    source.write_bytes(b"fake-hwp")

    converted_dir = tmp_path / "converted"
    converted_dir.mkdir()
    converted = converted_dir / "sample.hwpx"

    xml = """<?xml version="1.0" encoding="UTF-8"?>
    <hs:sec xmlns:hs="urn:section" xmlns:hp="urn:para">
      <hp:p><hp:run><hp:t>한글 변환 성공</hp:t></hp:run></hp:p>
    </hs:sec>
    """
    with zipfile.ZipFile(converted, "w") as archive:
        archive.writestr("Contents/section0.xml", xml)

    monkeypatch.setattr(
        "app.document_loader.convert_hwp_to_hwpx",
        lambda _path: converted,
    )

    text = load_document(source)

    assert "한글 변환 성공" in text
    assert not converted_dir.exists()


def test_hwpx_exposes_structured_table_metadata(tmp_path: Path):
    path = tmp_path / "structured.hwpx"
    xml = """<?xml version="1.0" encoding="UTF-8"?>
    <hs:sec xmlns:hs="urn:section" xmlns:hp="urn:para">
      <hp:p>
        <hp:run>
          <hp:tbl colCnt="3" rowCnt="2">
            <hp:tr>
              <hp:tc>
                <hp:subList><hp:p><hp:run><hp:t>항목</hp:t></hp:run></hp:p></hp:subList>
                <hp:cellAddr colAddr="0" rowAddr="0"/>
                <hp:cellSpan colSpan="1" rowSpan="1"/>
              </hp:tc>
              <hp:tc>
                <hp:subList><hp:p><hp:run><hp:t>세부 내용</hp:t></hp:run></hp:p></hp:subList>
                <hp:cellAddr colAddr="1" rowAddr="0"/>
                <hp:cellSpan colSpan="2" rowSpan="1"/>
              </hp:tc>
            </hp:tr>
            <hp:tr>
              <hp:tc>
                <hp:subList><hp:p><hp:run><hp:t>1</hp:t></hp:run></hp:p></hp:subList>
                <hp:cellAddr colAddr="0" rowAddr="1"/>
              </hp:tc>
              <hp:tc>
                <hp:subList>
                  <hp:p><hp:run><hp:t>첫째 문단</hp:t></hp:run></hp:p>
                  <hp:p><hp:run><hp:t>둘째 문단</hp:t></hp:run></hp:p>
                </hp:subList>
                <hp:cellAddr colAddr="1" rowAddr="1"/>
                <hp:cellSpan colSpan="2" rowSpan="1"/>
              </hp:tc>
            </hp:tr>
          </hp:tbl>
        </hp:run>
      </hp:p>
    </hs:sec>
    """

    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("Contents/section0.xml", xml)

    content = load_document_content(path)

    assert content.table_structure_supported is True
    assert len(content.tables) == 1
    table = content.tables[0]
    assert table.rows == 2
    assert table.cols == 3
    assert table.merged_cells == 2
    assert table.cells[1].col_span == 2
    assert any(cell.paragraph_count >= 2 for cell in table.cells)


def test_docx_exposes_structured_table_metadata(tmp_path: Path):
    path = tmp_path / "structured.docx"
    doc = Document()
    doc.add_paragraph("표 앞")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "운영 내용"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "내용 A"
    table.cell(1, 2).text = "내용 B"
    doc.save(path)

    content = load_document_content(path)

    assert content.table_structure_supported is True
    assert len(content.tables) == 1
    extracted = content.tables[0]
    assert extracted.rows == 2
    assert extracted.cols == 3
    assert extracted.merged_cells >= 1
    assert "운영 내용" in content.text
