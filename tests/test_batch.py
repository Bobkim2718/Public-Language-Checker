from pathlib import Path

from app.analyzer import PublicLanguageAnalyzer
from app.batch import MAX_BATCH_FILES, analyze_files, select_batch_files


RULES = {
    "weights": {
        "알기 쉬운 용어": 35,
        "알기 쉬운 문장": 35,
        "어문규범": 20,
        "한글 사용": 10,
    },
    "sentence": {"recommended_max_chars": 100, "severe_max_chars": 150},
}


def test_select_batch_files_caps_at_20_and_ignores_duplicates(tmp_path: Path):
    paths = []
    for index in range(25):
        path = tmp_path / f"{index:02d}.txt"
        path.write_text("문서", encoding="utf-8")
        paths.append(path)

    selected = select_batch_files(paths + [paths[0]])

    assert len(selected) == MAX_BATCH_FILES
    assert len({str(x) for x in selected}) == MAX_BATCH_FILES


def test_analyze_multiple_text_files(tmp_path: Path):
    first = tmp_path / "a.txt"
    second = tmp_path / "b.txt"
    first.write_text("학생 의견을 반영한다.", encoding="utf-8")
    second.write_text("사업을 모니터링한다.", encoding="utf-8")

    analyzer = PublicLanguageAnalyzer(
        [
            {
                "term": "모니터링",
                "alternatives": ["점검"],
                "alt_text": "점검",
                "severity": "review",
                "source": "official",
                "source_type": "OFFICIAL",
            }
        ],
        RULES,
    )

    results = analyze_files([first, second], analyzer)

    assert len(results) == 2
    assert all(item.ok for item in results)
    assert results[0].result.total_score == 100
    assert results[1].result.total_score < 100
    assert results[1].result.stats["official_unique"] == 1


def test_batch_includes_table_readability_for_docx(tmp_path: Path):
    from docx import Document

    path = tmp_path / "table.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "내용"
    table.cell(1, 0).text = "대상"
    table.cell(1, 1).text = "학생"
    doc.save(path)

    analyzer = PublicLanguageAnalyzer([], RULES)
    results = analyze_files([path], analyzer, rules=RULES)

    assert results[0].ok
    assert results[0].table_result is not None
    assert results[0].table_result.supported is True
    assert results[0].table_result.table_count == 1
